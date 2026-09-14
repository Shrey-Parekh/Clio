"""Reading documents and projects (6.9), against real files.

Every file here is genuinely written to disk and genuinely parsed - a real
.docx, .pptx, .xlsx and .pdf, built in a temp folder. Mocking the libraries
would only prove that the mocks agree with each other, and the thing most likely
to break is a real file in a format nobody tested.

Run: python tests/test_documents.py
"""

import asyncio
import json
import shutil
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from clio.capabilities.files import Request, lookup  # noqa: E402
from clio.core import documents  # noqa: E402
from clio.llm import longform  # noqa: E402


class FakeLLM:
    def __init__(self):
        self.calls = 0
        self.tiers = []

    async def complete(self, messages, tier="default"):
        self.calls += 1
        self.tiers.append(tier)
        return f"summary {self.calls}"


def make_pdf(path: Path, text: str) -> None:
    """A real PDF, assembled by hand with correct cross-reference offsets. The
    alternative is another dependency solely to write one test fixture."""
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 200] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        None,  # the content stream, built below
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    stream = f"BT /F1 12 Tf 20 100 Td ({text}) Tj ET".encode()
    objects[3] = (b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n"
                  + stream + b"\nendstream")

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"

    start = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{start}\n%%EOF\n").encode()
    path.write_bytes(bytes(out))


async def main():
    root = Path(tempfile.mkdtemp(prefix="clio-docs-"))
    try:
        # --- Word ---

        import docx

        word = docx.Document()
        word.add_paragraph("Enrolment brief for the autumn term.")
        word.add_paragraph("Submit 2000 words by the 30th of September.")
        table = word.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Criterion"
        table.cell(0, 1).text = "Weight"
        table.cell(1, 0).text = "Analysis"
        table.cell(1, 1).text = "40 percent"
        word.save(str(root / "brief.docx"))

        found = documents.extract(root / "brief.docx")
        assert found.kind == "word"
        assert "2000 words" in found.text and "Criterion | Weight" in found.text, found.text
        assert "Word document" in found.summary and "1 table" in found.summary, found.summary

        # --- PowerPoint, including the speaker notes ---

        from pptx import Presentation

        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = "E-waste training"
        slide.placeholders[1].text = "Sorting, dismantling, recovery"
        slide.notes_slide.notes_text_frame.text = "Mention the recovery rate figures."
        deck.save(str(root / "deck.pptx"))

        found = documents.extract(root / "deck.pptx")
        assert found.kind == "slides" and "deck of 1 slides" in found.summary
        assert "E-waste training" in found.text
        assert "(notes) Mention the recovery rate" in found.text, "notes carry the actual argument"

        # --- Excel, two sheets ---

        from openpyxl import Workbook

        book = Workbook()
        first = book.active
        first.title = "Marks"
        first.append(["Student", "Mark"])
        for i in range(20):
            first.append([f"Student {i}", 40 + i])
        second = book.create_sheet("Notes")
        second.append(["Comment"])
        book.save(str(root / "marks.xlsx"))

        found = documents.extract(root / "marks.xlsx")
        assert found.kind == "sheet"
        assert "Marks (21 rows, 2 columns)" in found.summary, found.summary
        assert "Notes (1 rows, 1 columns)" in found.summary, found.summary
        assert "Student | Mark" in found.text
        print("OK  Word, PowerPoint and Excel read, including tables and speaker notes")

        # --- PDF, and a scan with no text in it ---

        make_pdf(root / "paper.pdf", "Recovery rates rose to 62 percent")
        found = documents.extract(root / "paper.pdf")
        assert found.kind == "pdf", found
        assert "Recovery rates rose" in found.text, found.text
        assert "1-page PDF" in found.summary, found.summary

        make_pdf(root / "scan.pdf", "")
        scanned = documents.extract(root / "scan.pdf")
        assert scanned.text.strip() == "" and "scan" in scanned.summary, scanned.summary
        print(f"OK  PDF text extracted, and a scan says so: {scanned.summary!r}")

        # --- tables and data, on the standard library ---

        (root / "grades.csv").write_text(
            "name,mark,grade\n" + "\n".join(f"student {i},{50+i},B" for i in range(30)),
            encoding="utf-8")
        found = documents.extract(root / "grades.csv")
        assert found.kind == "table"
        assert "30 rows and 3 columns" in found.summary and "name, mark, grade" in found.summary
        assert found.truncated is True, "she says when she only read the first rows"

        # A semicolon export is still a csv to Windows.
        (root / "euro.csv").write_text("a;b;c\n1;2;3\n", encoding="utf-8")
        assert "3 columns" in documents.extract(root / "euro.csv").summary

        (root / "config.json").write_text(json.dumps(
            {"model": "gpt", "epochs": 5, "layers": [{"units": 64}, {"units": 32}]}),
            encoding="utf-8")
        found = documents.extract(root / "config.json")
        assert found.kind == "data"
        assert "3 keys" in found.summary, found.summary
        assert "a list of 2" in found.text, found.text
        print("OK  CSV sniffed and counted, JSON described by shape rather than dumped")

        # --- images and unknown formats are refused honestly ---

        (root / "photo.jpg").write_bytes(b"\xff\xd8\xff\xe0 not really a jpeg")
        image = documents.extract(root / "photo.jpg")
        assert image.kind == "image" and image.text == ""
        assert "no way to look at pictures" in image.summary
        (root / "archive.zip").write_bytes(b"PK\x03\x04")
        assert documents.extract(root / "archive.zip") is None

        # Found live, 2026-09-14: three of these were sitting in his Documents.
        # Word, Excel and PowerPoint leave them while a file is open. They carry
        # the right extension and none of the contents, so every parser throws.
        (root / "~$brief.docx").write_bytes(b"not a docx at all")
        lock = documents.extract(root / "~$brief.docx")
        assert lock.kind == "lock" and "while a document is open" in lock.summary, lock
        print("OK  an image says why it can't be read; Office lock files aren't parsed")

        # --- projects: what it is, and the command it runs with ---

        python_project = root / "ewaste"
        (python_project / "ewaste").mkdir(parents=True)
        (python_project / "README.md").write_text(
            "# E-waste training\n\nTrains the sorting model.", encoding="utf-8")
        (python_project / "pyproject.toml").write_text(
            "[project]\nname='ewaste'\n", encoding="utf-8")
        (python_project / "ewaste" / "__main__.py").write_text("print('training')", encoding="utf-8")

        project = documents.describe_project(python_project)
        assert project.kind == "project"
        assert "a Python project" in project.summary, project.summary
        assert "run with python -m ewaste" in project.summary, project.summary
        assert "Trains the sorting model" in project.text

        node = root / "site"
        node.mkdir()
        (node / "package.json").write_text(json.dumps(
            {"name": "site", "scripts": {"dev": "vite", "build": "vite build"}}), encoding="utf-8")
        assert documents.run_command(node) == "npm run dev"

        loose = root / "scripts"
        loose.mkdir()
        (loose / "train.py").write_text("print('go')", encoding="utf-8")
        assert documents.run_command(loose) == "python train.py"

        (root / "make-it").mkdir()
        (root / "make-it" / "Makefile").write_text("all:\n\techo hi\n", encoding="utf-8")
        assert documents.run_command(root / "make-it") == "make"

        # Found live, 2026-09-14: the repo folder is "Clio" and the package is
        # "clio". Windows matches paths case-insensitively, so looking for
        # `folder / folder.name` produced "python -m Clio" - right here, wrong
        # everywhere else, and it is the command 7.2 would actually run.
        # (And this folder is deliberately not named "Ewaste": on Windows that
        # collides with the "ewaste" project above, which is the same bug again.)
        cased = root / "Sorter"
        (cased / "sorter").mkdir(parents=True)
        (cased / "pyproject.toml").write_text("[project]\n", encoding="utf-8")
        (cased / "sorter" / "__main__.py").write_text("print('go')", encoding="utf-8")
        assert documents.run_command(cased) == "python -m sorter", documents.run_command(cased)

        empty = root / "nothing"
        empty.mkdir()
        assert documents.run_command(empty) == "", "a folder with no manifest gets no invented command"
        assert "folder of files" in documents.describe_project(empty).summary
        print(f"OK  projects understood: {project.summary}")

        # --- long documents are chunked, not truncated ---

        assert len(longform.split("a" * 20_000, chunk_chars=6_000)) == 4
        paragraphs = "\n\n".join(f"Paragraph {i}. " + "word " * 200 for i in range(20))
        chunks = longform.split(paragraphs, chunk_chars=6_000)
        assert all(len(c) <= 6_500 for c in chunks), [len(c) for c in chunks]

        llm = FakeLLM()
        await longform.summarise("short document", "what is it", "p", llm)
        assert llm.calls == 1, "a short document is one call, not a chunking ceremony"

        llm = FakeLLM()
        spoken = await longform.summarise(paragraphs, "what is it", "p", llm,
                                          max_chunks=3, chunk_chars=6_000, pace_s=0)
        assert llm.calls == 4, llm.calls          # three parts, then the whole
        assert llm.tiers[:3] == ["fast", "fast", "fast"], llm.tiers
        assert "percent of it" in spoken, spoken
        print(f"OK  long documents chunked and paced, and the cap is admitted: {spoken!r}")

        # --- through the file capability ---

        spoken, to_summarise = lookup(Request("read", "", (root / "marks.xlsx",)), (root,))
        assert "spreadsheet" in spoken and to_summarise == ""

        spoken, to_summarise = lookup(Request("summarise", "", (root / "brief.docx",)), (root,))
        assert spoken == "" and "2000 words" in to_summarise, "summarising hands the text to the model"

        spoken, _ = lookup(Request("read", "", (root / "photo.jpg",)), (root,))
        assert "no way to look at pictures" in spoken

        spoken, _ = lookup(Request("read", "", (python_project,)), (root,))
        assert "run with python -m ewaste" in spoken, spoken

        spoken, to_summarise = lookup(Request("summarise", "", (python_project,)), (root,))
        assert spoken == "" and "Trains the sorting model" in to_summarise
        print("OK  the existing 'read that file' path now reads documents and projects too")

        print("\nAll document checks passed.")
    finally:
        shutil.rmtree(root, ignore_errors=True)


if __name__ == "__main__":
    asyncio.run(main())
